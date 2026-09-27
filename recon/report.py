from __future__ import annotations
import html
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from urllib.parse import urlsplit
from .model import priority, redact_url, SECRET_KEY


def clean(value):
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return value
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value))
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) else text


def snapshot(db, run):
    evidence = defaultdict(list)
    for e in db.rows("evidence"):
        evidence[e["endpoint_id"]].append({"source": e["source"], "source_url": redact_url(e["source_url"]) if e["source_url"] else "", **json.loads(e["data"])})
    obs = {r["endpoint_id"]: r for r in db.rows("observations") if r["run_id"] == run}
    changes = [r for r in db.rows("changes") if r["run_id"] == run]
    changed = {r["endpoint_id"] for r in changes}
    endpoints = []
    for row in db.rows("endpoints"):
        row.update(obs.get(row["id"], {"state": "not_probed", "status": None}))
        row["evidence"] = evidence[row["id"]]
        row["priority"] = priority(row, row["id"] in changed)
        row["url"] = redact_url(row["url"])
        row["pattern"] = redact_url(row["pattern"])
        if (row.get("location") or "").startswith(("http://", "https://")):
            row["location"] = redact_url(row["location"])
        endpoints.append(row)
    for c in changes:
        if c["kind"] == "new_endpoint":
            c["after_value"] = redact_url(c["after_value"])
        if c["kind"] == "location_changed":
            for f in ("before_value", "after_value"):
                if c[f].startswith(("http://", "https://")):
                    c[f] = redact_url(c[f])
    findings = db.rows("findings")
    for f in findings:
        f["source_url"] = redact_url(f["source_url"])
    clusters = defaultdict(list)
    for e in endpoints:
        if e.get("body_hash"):
            clusters[e["body_hash"]].append(e["id"])
    return {"run": run, "endpoints": sorted(endpoints, key=lambda x: (-x["priority"]["score"], x["url"], x["method"])),
            "changes": changes, "findings": findings, "assets": db.rows("assets"),
            "events": [r for r in db.rows("events") if r["run_id"] == run],
            "response_clusters": {k: v for k, v in clusters.items() if len(v) > 1}}


def export(db, run, output):
    out = Path(output)
    out.mkdir(parents=True, exist_ok=True)
    data = snapshot(db, run)
    (out / "report.json").write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    collection = {"info": {"name": "Passive Scan v9 — manual review", "schema": "https://schema.getpostman.com/json/collection/v2.1.0/collection.json"}, "item": []}
    for e in data["endpoints"]:
        if e["method"] not in {"GET", "POST", "PUT", "PATCH", "DELETE", "HEAD", "OPTIONS", "TRACE"}:
            continue  # unknown methods stay in the report, never fabricated as GET
        fields = sorted({p["name"] for ev in e["evidence"] for p in ev.get("parameters", []) if p.get("in") == "body"})
        request = {"method": e["method"], "url": e["url"], "header": [],
                   "description": "Imported structure only. Supply authorized credentials and real body values manually. Nested fields are flattened placeholders."}
        if fields:
            request["body"] = {"mode": "raw", "raw": json.dumps({f: "{{" + re.sub(r'\W', '_', f) + "}}" for f in fields}), "options": {"raw": {"language": "json"}}}
        collection["item"].append({"name": e["method"] + " " + urlsplit(e["url"]).path, "request": request})
    (out / "postman.json").write_text(json.dumps(collection, ensure_ascii=False, indent=2), encoding="utf-8")
    tables = {
        "Endpoints": (["ID", "Score", "Method", "URL", "State", "HTTP", "Tags", "Sources", "Review", "Owner", "Note", "Next step"],
                      [[e["id"], e["priority"]["score"], e["method"], e["url"], e["state"], e.get("status"), ", ".join(e["priority"]["tags"]), ", ".join(sorted({x["source"] for x in e["evidence"]})), e["review"], e["owner"], e["note"], e["priority"]["next_step"]] for e in data["endpoints"]]),
        "Changes": (["Kind", "Endpoint ID", "Before", "After", "Confirmed"], [[c["kind"], c["endpoint_id"], c["before_value"], c["after_value"], c["confirmed"]] for c in data["changes"]]),
        "Secrets": (["Detector", "Source", "Location", "Value", "Verification", "Review"], [[f["detector"], f["source_url"], f["location"], f["masked"], f["verification"], f["review"]] for f in data["findings"]]),
        "Assets": (["Host", "First run", "Last run"], [[x["host"], x["first_run"], x["last_run"]] for x in data["assets"]]),
        "Pipeline": (["Stage", "State", "Detail"], [[x["stage"], x["state"], x["detail"]] for x in data["events"]]),
    }
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    for name, (headers, rows) in tables.items():
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append([clean(v) for v in row])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="18324F")
            cell.font = Font(color="FFFFFF", bold=True)
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = 65 if h in {"URL", "Source", "Next step", "Before", "After", "Note"} else 22
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(out / "report.xlsx")
    esc = lambda v: html.escape(str(v if v is not None else ""), quote=True)
    sections = []
    for name, (headers, rows) in tables.items():
        sections.append("<section><h2>" + esc(name) + "</h2><div class='scroll'><table><thead><tr>" + "".join("<th>" + esc(h) + "</th>" for h in headers) + "</tr></thead><tbody>" + "".join("<tr>" + "".join("<td>" + esc(v) + "</td>" for v in row) + "</tr>" for row in rows) + "</tbody></table></div></section>")
    states = Counter(e["state"] for e in data["endpoints"])
    page = """<!doctype html><html lang="ko"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
    <meta http-equiv="Content-Security-Policy" content="default-src 'none'; style-src 'unsafe-inline'; img-src 'self';">
    <title>Passive Scan v9</title><style>body{font:15px system-ui;background:#f2f5fa;color:#18324f;margin:24px}h1{font-size:30px}section{background:white;padding:20px;margin:20px 0;border-radius:12px}.scroll{overflow:auto}table{border-collapse:collapse;width:100%}td,th{padding:10px;border-bottom:1px solid #ddd;text-align:left;max-width:560px;overflow-wrap:anywhere}th{background:#18324f;color:white}tr:nth-child(even){background:#f6f8fb}p{line-height:1.7}</style>"""
    page += f"<h1>Passive Scan v9 · Run {run}</h1><p>우선순위 점수는 취약점 확률이 아닙니다. 변경 목록은 추가 확인이 필요한 관찰입니다.</p><p>{esc(dict(states))} · 동일 응답 그룹 {len(data['response_clusters'])}개</p>" + "".join(sections) + "</html>"
    (out / "index.html").write_text(page, encoding="utf-8")
    return data
